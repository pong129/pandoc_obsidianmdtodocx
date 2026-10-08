from pathlib import Path
import subprocess
import zipfile

from lxml import etree as ET
import pytest

from experiments.tibetan_typography.prototype import (
    A, W, Typography, is_tibetan, process_package, segments, split_run,
)

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
CONFIG = Typography()


@pytest.mark.parametrize("text", [
    "한국어Englishབོད་ཡིག་끝", "བསྒྲུབས་ སྐྱེས་ ཨོཾ་ ཧཱུྃ།",
    "༡༢༣།123", "(བོད་),한국어!", "sangs rgyas Wylie", "",
    "བོད\u200d་Latin", "emoji😀 བོད་", "e\u0301한글བོད་",
])
def test_segments_preserve_text_and_classify(text):
    result = segments(text)
    assert "".join(value for _, value in result) == text
    for tibetan, value in result:
        for char in value:
            if char not in "\u200c\u200d":
                assert is_tibetan(char) == tibetan


@pytest.mark.parametrize("ns", [W, A])
def test_split_preserves_style_link_and_nontext(ns):
    p = ET.Element(f"{{{ns}}}p")
    r = ET.SubElement(p, f"{{{ns}}}r")
    props = ET.SubElement(r, f"{{{ns}}}rPr")
    if ns == W:
        ET.SubElement(props, f"{{{ns}}}b")
        ET.SubElement(props, f"{{{ns}}}i")
    else:
        props.set("b", "1"); props.set("i", "1")
        ET.SubElement(props, f"{{{ns}}}hlinkClick", {"id": "keep-link"})
    ET.SubElement(r, f"{{{ns}}}t").text = "한글ABCབོད་DEF"
    ET.SubElement(r, f"{{{ns}}}tab")
    ET.SubElement(r, f"{{{ns}}}br")
    ET.SubElement(r, f"{{{ns}}}t").text = "끝ཆོས་"
    before = "".join(p.itertext())
    split_run(r, ns, 12, CONFIG)
    assert "".join(p.itertext()) == before
    assert len(p.findall(f".//{{{ns}}}tab")) == 1
    assert len(p.findall(f".//{{{ns}}}br")) == 1
    for new_run in p:
        new_props = new_run.find(f"{{{ns}}}rPr")
        if ns == W:
            assert new_props.find(f"{{{ns}}}b") is not None
            assert new_props.find(f"{{{ns}}}i") is not None
        else:
            assert new_props.get("b") == "1" and new_props.get("i") == "1"
            assert new_props.find(f"{{{ns}}}hlinkClick").get("id") == "keep-link"


def read_xml(path):
    with zipfile.ZipFile(path) as archive:
        return {name: ET.fromstring(archive.read(name)) for name in archive.namelist()
                if name.endswith(".xml") and not name.endswith("[Content_Types].xml")}


def assert_typography(run, ns, config):
    text = "".join(node.text or "" for node in run.findall(f"{{{ns}}}t"))
    if not text:
        return None
    flags = {is_tibetan(char) for char in text}
    assert len(flags) == 1, f"Mixed run remains: {text!r}"
    tibetan = flags.pop()
    props = run.find(f"{{{ns}}}rPr")
    if ns == W:
        fonts = props.find(f"{{{W}}}rFonts")
        assert fonts.get(f"{{{W}}}ascii") == (config.tibetan_font if tibetan else config.latin_font)
        assert fonts.get(f"{{{W}}}eastAsia") == (config.tibetan_font if tibetan else config.korean_font)
        size = int(props.find(f"{{{W}}}sz").get(f"{{{W}}}val")) / 2
    else:
        assert props.find(f"{{{A}}}latin").get("typeface") == (config.tibetan_font if tibetan else config.latin_font)
        assert props.find(f"{{{A}}}ea").get("typeface") == (config.tibetan_font if tibetan else config.korean_font)
        size = int(props.get("sz")) / 100
    assert size in ((config.tibetan_pt, config.tibetan_size(config.heading_pt)) if tibetan
                    else (config.body_pt, config.heading_pt))
    return tibetan


@pytest.mark.parametrize("kind", ["docx", "pptx"])
@pytest.mark.parametrize("config", [
    CONFIG,
    Typography(latin_font="Times New Roman", korean_font="Malgun Gothic",
               tibetan_font="Monlam Bodyig", body_pt=15, tibetan_pt=27, heading_pt=30),
])
def test_pandoc_conversion_and_exact_typography(tmp_path, kind, config):
    original, output = tmp_path / f"original.{kind}", tmp_path / f"result.{kind}"
    subprocess.run([str(ROOT / "pandoc.exe"), str(HERE / "sample.md"), "-o", str(original)],
                   check=True, capture_output=True, timeout=30)
    process_package(original, output, config, kind)
    before, after = read_xml(original), read_xml(output)
    ns = W if kind == "docx" else A
    counts = {True: 0, False: 0}
    for name, tree in before.items():
        assert "".join(tree.itertext()) == "".join(after[name].itertext()), name
        # No lost hyperlinks, line breaks, pictures or footnote references.
        for tag in ("hyperlink", "drawing", "blip", "br", "tab", "footnoteReference"):
            assert len(tree.findall(f".//{{*}}{tag}")) == len(after[name].findall(f".//{{*}}{tag}"))
        # DrawingML links live on runs, so splitting legitimately duplicates the
        # property. Relationship targets must stay identical, with none lost.
        for tag in ("hlinkClick", "hlinkMouseOver"):
            assert {tuple(sorted(n.attrib.items())) for n in tree.findall(f".//{{*}}{tag}")} == {
                tuple(sorted(n.attrib.items())) for n in after[name].findall(f".//{{*}}{tag}")}
        relevant = name.startswith("word/") if kind == "docx" else name.startswith("ppt/slides/slide")
        if relevant:
            for run in after[name].iter(f"{{{ns}}}r"):
                flag = assert_typography(run, ns, config)
                if flag is not None:
                    counts[flag] += 1
    assert counts[True] >= 10 and counts[False] >= 10, counts
    if kind == "docx":
        assert any(is_tibetan(c) for c in "".join(after['word/footnotes.xml'].itertext()))
    # Applying the algorithm twice must not keep splitting or alter text/styles.
    second = tmp_path / f"second.{kind}"
    process_package(output, second, config, kind)
    assert {k: ET.tostring(v) for k, v in after.items()} == {k: ET.tostring(v) for k, v in read_xml(second).items()}


def test_docx_header_and_table(tmp_path):
    from docx import Document
    doc = Document()
    doc.sections[0].header.paragraphs[0].add_run("Header한글བོད་")
    doc.sections[0].footer.paragraphs[0].add_run("Footerཆོས་한글")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table한글བོད་"
    original, result = tmp_path/'source.docx', tmp_path/'result.docx'
    doc.save(original); process_package(original, result, CONFIG, 'docx')
    trees = read_xml(result)
    for name in ('word/header1.xml', 'word/footer1.xml', 'word/document.xml'):
        assert any(assert_typography(r, W, CONFIG) for r in trees[name].iter(f'{{{W}}}r'))


def test_pptx_table_and_bold_italic_hyperlink(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(7), Inches(2)).table
    r = table.cell(0, 0).text_frame.paragraphs[0].add_run()
    r.text = "Table한글བོད་ABC"; r.font.bold = True; r.font.italic = True
    r.hyperlink.address = 'https://example.com'
    original, result = tmp_path/'source.pptx', tmp_path/'result.pptx'
    prs.save(original); process_package(original, result, CONFIG, 'pptx')
    edited = Presentation(result)
    runs = edited.slides[0].shapes[0].table.cell(0, 0).text_frame.paragraphs[0].runs
    assert ''.join(r.text for r in runs) == 'Table한글བོད་ABC'
    assert len(runs) == 3
    assert all(r.font.bold and r.font.italic and r.hyperlink.address == 'https://example.com' for r in runs)
    for r in read_xml(result)['ppt/slides/slide1.xml'].iter(f'{{{A}}}r'):
        assert_typography(r, A, CONFIG)


def test_reproduce_current_docx_mixed_run_bug(tmp_path):
    from docx import Document
    from experiments.tibetan_typography.legacy_docx import PostProcessor
    doc = Document(); doc.add_paragraph('한국어ABCབོད་DEF')
    path = tmp_path/'legacy.docx'; doc.save(path)
    PostProcessor.process(str(path), {
        'font_ascii': 'Arial', 'font_korean': 'Malgun Gothic',
        'font_tibetan': CONFIG.tibetan_font, 'size_body': 12, 'size_cs': 20,
    })
    run = next(read_xml(path)['word/document.xml'].iter(f'{{{W}}}r'))
    # Reproduce the old failure: Latin and Korean in this run also become
    # Tibetan font at 20pt, despite the body size being 12pt.
    fonts = run.find(f'{{{W}}}rPr/{{{W}}}rFonts')
    assert fonts.get(f'{{{W}}}ascii') == CONFIG.tibetan_font
    assert fonts.get(f'{{{W}}}eastAsia') == CONFIG.tibetan_font
    assert run.find(f'{{{W}}}rPr/{{{W}}}sz').get(f'{{{W}}}val') == '40'
    result = tmp_path/'fixed.docx'; process_package(path, result, CONFIG, 'docx')
    for r in read_xml(result)['word/document.xml'].iter(f'{{{W}}}r'):
        assert_typography(r, W, CONFIG)


def test_reproduce_current_pptx_tibetan_setting_ignored(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from experiments.tibetan_typography.legacy_pptx import PptxPostProcessor
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    box.text = '한국어ABCབོད་DEF'
    path = tmp_path/'legacy.pptx'; prs.save(path)
    PptxPostProcessor.process(str(path), {
        'size_body': 12, 'size_cs': 20, 'font_tibetan': CONFIG.tibetan_font,
    })
    run = next(read_xml(path)['ppt/slides/slide1.xml'].iter(f'{{{A}}}r'))
    props = run.find(f'{{{A}}}rPr')
    assert props.get('sz') == '1200'  # Tibetan size ignored: still 12pt.
    assert props.find(f'{{{A}}}cs') is None  # Tibetan font also ignored.
    result = tmp_path/'fixed.pptx'; process_package(path, result, CONFIG, 'pptx')
    for r in read_xml(result)['ppt/slides/slide1.xml'].iter(f'{{{A}}}r'):
        assert_typography(r, A, CONFIG)
