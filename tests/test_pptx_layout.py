from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches
from src.core.pptx_converter import PptxConverter
from src.core.pandoc_manager import PandocManager
from src.core.pptx_post_processor import PptxPostProcessor


def test_untitled_pandoc_slide_reclaims_empty_title_band(tmp_path):
    source = tmp_path/'untitled.md'
    source.write_text('བོད་ཡིག་ སྐྱེས་\n\n한국어 English\n', encoding='utf-8')
    config = {'size_body':20, 'size_cs':32, 'insert_title':False}
    result = PptxConverter(PandocManager()).convert([str(source)], str(tmp_path), config)[0]
    prs = Presentation(result)
    body = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text)
    assert body.top == Inches(.25)
    assert body.top + body.height == prs.slide_height - Inches(.25)
    assert body.text_frame.margin_top == 0
    assert '한국어 English' in body.text
    assert any(r.font.size.pt == 32 for p in body.text_frame.paragraphs for r in p.runs if 'བོད' in r.text)


def test_real_title_keeps_existing_layout(tmp_path):
    source = tmp_path/'titled.md'; source.write_text('## 제목\n\n본문\n', encoding='utf-8')
    result = PptxConverter(PandocManager()).convert([str(source)], str(tmp_path), {'insert_title':False})[0]
    prs = Presentation(result)
    body = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text == '본문')
    assert body.top > Inches(1)
    assert prs.slides[0].shapes.title.text == '제목'


def test_side_by_side_columns_move_together_and_expand(tmp_path):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[3])
    slide.shapes.title.text = ''
    bodies = [s for s in slide.placeholders if s != slide.shapes.title]
    for i, body in enumerate(bodies): body.text = 'Column ' + str(i)
    original_left = [b.left for b in bodies]
    path = tmp_path/'columns.pptx'; prs.save(path)
    PptxPostProcessor.process(str(path), {'size_body':20, 'size_cs':32})
    edited = Presentation(path)
    bodies = [s for s in edited.slides[0].shapes if s.has_text_frame and s.text]
    assert [b.left for b in bodies] == original_left
    assert all(b.top == Inches(.25) for b in bodies)
    assert all(b.top + b.height == edited.slide_height - Inches(.25) for b in bodies)


def test_body_does_not_expand_over_table_below(tmp_path):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = ''
    body = slide.placeholders[1]; body.text = 'Textབོད་'
    body.top = Inches(1.5); body.height = Inches(1)
    table = slide.shapes.add_table(1,1, Inches(.5),Inches(3), Inches(9), Inches(1))
    table.table.cell(0,0).text = 'Table'
    original_gap = table.top - body.top
    path=tmp_path/'table.pptx';prs.save(path)
    PptxPostProcessor.process(str(path), {'size_body':20})
    edited = Presentation(path); body = edited.slides[0].placeholders[1]
    table = next(s for s in edited.slides[0].shapes if s.has_table)
    assert body.top == Inches(.25)
    assert body.height == Inches(1)
    assert table.top - body.top == original_gap
