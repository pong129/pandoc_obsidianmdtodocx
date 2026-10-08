from pathlib import Path
import zipfile
from lxml import etree as ET
from pptx import Presentation
from pptx.util import Inches
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QFontDatabase
import pytest
from src.core.pptx_converter import PptxConverter
from src.core.pandoc_manager import PandocManager
from src.core.pptx_pagination import paginate_package, paragraph_text, A, P
from src.core.pptx_post_processor import PptxPostProcessor


@pytest.fixture(scope='module', autouse=True)
def app():
    application = QApplication.instance() or QApplication([])
    for name in ('malgun.ttf', 'arial.ttf', 'himalaya.ttf'):
        QFontDatabase.addApplicationFont('C:/Windows/Fonts/' + name)
    return application


def body_text(path):
    return ''.join(s.text for slide in Presentation(path).slides for s in slide.shapes
                   if s.has_text_frame and s != slide.shapes.title)


def test_long_mixed_paragraph_preserves_text_sizes_and_hyperlink(tmp_path):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = '제목'
    body = slide.placeholders[1]
    paragraph = body.text_frame.paragraphs[0]
    expected = ('😀 e\u0301 བསྒྲུབས་སྐྱེས་ English ' * 200)
    run = paragraph.add_run(); run.text = expected; run.font.bold = True
    run.hyperlink.address = 'https://example.org/source'
    path = tmp_path/'mixed.pptx'; prs.save(path)
    PptxPostProcessor.process(str(path), {'size_body':20, 'size_cs':20, 'size_east_asia':18})
    assert paginate_package(path) > 1
    edited = Presentation(path)
    assert body_text(path) == expected
    for index, page in enumerate(edited.slides):
        assert page.shapes.title.text == ('제목' if index == 0 else '')
        runs = [r for p in page.placeholders[1].text_frame.paragraphs for r in p.runs]
        assert all(r.font.bold and r.hyperlink.address == 'https://example.org/source' for r in runs)
        assert all(r.font.size.pt in (18,20) for r in runs)
        first = next((r.text[0] for r in runs if r.text), '')
        assert not ('\u0f90' <= first <= '\u0fbc')
    with zipfile.ZipFile(path) as archive:
        types = ET.fromstring(archive.read('[Content_Types].xml'))
        names = {n.get('PartName') for n in types}
        for n in archive.namelist():
            if n.startswith('ppt/slides/slide') and n.endswith('.xml'):
                assert '/' + n in names


def test_manual_break_preserved_and_auto_split_can_be_disabled(tmp_path):
    source = tmp_path/'source.md'
    source.write_text(('བོད་ཡིག་ ' * 20 + '\n한국어 번역 ' * 3 + '\n\n') * 25 + '\n\n***\n\n마지막 슬라이드', encoding='utf-8')
    converter = PptxConverter(PandocManager())
    base = {'insert_title':False, 'preserve_breaks':True}
    manual = converter.convert([str(source)], str(tmp_path), dict(base, auto_paginate=False))[0]
    auto = converter.convert([str(source)], str(tmp_path), dict(base, auto_paginate=True))[0]
    assert len(Presentation(manual).slides) == 2
    assert len(Presentation(auto).slides) > 2
    assert body_text(auto).replace('\n','') == body_text(manual).replace('\n','')
    assert any('마지막 슬라이드' in s.text for s in Presentation(auto).slides[-1].shapes if s.has_text_frame)


def test_hard_breaks_and_multi_paragraph_text_are_preserved(tmp_path):
    source = tmp_path/'breaks.md'
    content = '\n\n'.join(('བོད་ཡིག་ ' * 15) + '\n한국어 번역 ' * 3 for _ in range(8))
    source.write_text(content, encoding='utf-8')
    converter = PptxConverter(PandocManager())
    config = {'insert_title':False,'preserve_breaks':True,'auto_paginate':False}
    original = converter.convert([str(source)], str(tmp_path), config)[0]
    expected = body_text(original).replace('\n','').replace('\v','')
    paginate_package(original)
    assert body_text(original).replace('\n','').replace('\v','') == expected
    assert len(Presentation(original).slides) > 1


@pytest.mark.parametrize('separate_paragraph', [False, True])
def test_original_and_translation_always_move_together(tmp_path, separate_paragraph):
    source = tmp_path/'pairs.md'
    joiner = '\n\n' if separate_paragraph else '\n'
    source.write_text('\n\n'.join(f'T{i:03d} ' + 'བོད་ཡིག་ ' * 15 + joiner +
                                  f'K{i:03d} ' + '한국어 해석 ' * 10 for i in range(12)), encoding='utf-8')
    counts = []
    for size in (20, 28):
        converter = PptxConverter(PandocManager())
        output = converter.convert([str(source)], str(tmp_path), {
            'insert_title':False, 'size_cs':size, 'size_east_asia':18, 'preserve_breaks':True})[0]
        slides = Presentation(output).slides
        counts.append(len(slides))
        for i in range(12):
            containing = [page for page in slides if any(f'T{i:03d}' in s.text for s in page.shapes if s.has_text_frame)]
            assert len(containing) == 1
            assert any(f'K{i:03d}' in s.text for s in containing[0].shapes if s.has_text_frame)
        for page in slides:
            for shape in page.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if 'བོད' in run.text:
                                assert run.font.size.pt == size
    assert counts[1] >= counts[0]


def test_oversized_pair_alerts_without_splitting_or_shrinking(tmp_path):
    source = tmp_path/'oversized.md'
    source.write_text('བོད་ཡིག་ ' * 500 + '\n한국어 번역 ' * 500, encoding='utf-8')
    converter = PptxConverter(PandocManager())
    assert converter.convert([str(source)], str(tmp_path), {'insert_title':False,'preserve_breaks':True}) == []
    assert '한 묶음이 한 장보다' in converter.last_errors[0]
    assert 'བོད' in converter.last_errors[0]
    assert not list(tmp_path.glob('*.pptx'))


def test_long_title_expands_without_covering_body(tmp_path):
    source = tmp_path/'long_title.md'
    source.write_text('# ' + '여러 줄 제목 ' * 4 + '\n\nབོད་ཡིག་\n한국어 번역', encoding='utf-8')
    converter = PptxConverter(PandocManager())
    output = converter.convert([str(source)], str(tmp_path), {
        'insert_title':False,'preserve_breaks':True,'size_h1':44})[0]
    slide = Presentation(output).slides[0]
    title, body = slide.shapes.title, slide.placeholders[1]
    assert title.height > Inches(.94)
    assert title.top + title.height < body.top
