from pathlib import Path
import pytest
from pptx import Presentation
from PySide6.QtCore import QSettings
from src.core.pptx_converter import PptxConverter
from src.core.pandoc_manager import PandocManager
from src.core.settings_store import create_settings
from src.core.typography import from_config


@pytest.mark.parametrize('level', [1, 2, 3])
def test_section_heading_shares_slide_with_body(tmp_path, level):
    source = tmp_path / 'section.md'
    source.write_text('# 제1법륜\n\nབོད་ཡིག་\n\n한국어 본문\n', encoding='utf-8')
    output = PptxConverter(PandocManager()).convert(
        [str(source)], str(tmp_path), {'insert_title': False, 'slide_level': level})[0]
    slides = Presentation(output).slides
    assert len(slides) == 1
    assert slides[0].shapes.title.text == '제1법륜'
    text = '\n'.join(s.text for s in slides[0].shapes if s.has_text_frame)
    assert '한국어 본문' in text and 'བོད་ཡིག་' in text
    runs = [r for s in slides[0].shapes if s.has_text_frame
            for p in s.text_frame.paragraphs for r in p.runs]
    assert next(r.font.size.pt for r in runs if r.text == 'བོད་ཡིག་') == 20
    assert next(r.font.size.pt for r in runs if r.text.startswith('한국어')) == 18


def test_consecutive_headings_do_not_make_empty_section_slide(tmp_path):
    source = tmp_path / 'nested.md'
    source.write_text('# 장 제목\n\n## 소제목\n\n첫 본문\n\n## 다음\n\n다음 본문', encoding='utf-8')
    output = PptxConverter(PandocManager()).convert(
        [str(source)], str(tmp_path), {'insert_title': False})[0]
    slides = Presentation(output).slides
    assert len(slides) == 2
    first = '\n'.join(s.text for s in slides[0].shapes if s.has_text_frame)
    assert all(t in first for t in ('장 제목', '소제목', '첫 본문'))
    assert '{{H' not in first


def test_fenced_code_is_not_rewritten():
    converter = PptxConverter(PandocManager())
    for fence in ('```', '~~~~'):
        source = f'{fence}\n# code\n{fence}\n\n# 제목'
        result = converter.preprocess_content(source, {})
        assert f'{fence}\n# code\n{fence}' in result
        assert '## 제목 {{H1}}' in result


def test_previous_default_values_migrate_only_once(tmp_path):
    path = tmp_path / 'settings.ini'
    saved = QSettings(str(path), QSettings.IniFormat)
    saved.setValue('migration_complete', True)
    saved.setValue('pptx_size_cs', 32)
    saved.setValue('pptx_size_east_asia', 20)
    saved.setValue('docx_size_cs', 32)
    saved.setValue('presets/pptx/custom', '{"size_cs":32}')
    saved.sync()
    updated = create_settings(path)
    assert updated.value('pptx_size_cs', type=int) == 20
    assert updated.value('pptx_size_east_asia', type=int) == 18
    assert updated.value('docx_size_cs', type=int) == 32
    assert updated.value('presets/pptx/custom') == '{"size_cs":32}'
    updated.setValue('pptx_size_cs', 32)
    updated.sync()
    assert create_settings(path).value('pptx_size_cs', type=int) == 32


def test_custom_values_preserved(tmp_path):
    path = tmp_path / 'custom.ini'
    saved = QSettings(str(path), QSettings.IniFormat)
    saved.setValue('migration_complete', True)
    saved.setValue('pptx_size_cs', 25)
    saved.setValue('pptx_size_east_asia', 16)
    saved.sync()
    updated = create_settings(path)
    assert updated.value('pptx_size_cs', type=int) == 25
    assert updated.value('pptx_size_east_asia', type=int) == 16
    assert from_config({}, 'pptx').tibetan_pt == 20
    assert from_config({}, 'pptx').korean_pt == 18
